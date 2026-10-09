import csv
import io
import mimetypes
import secrets
import random
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4
from apscheduler.schedulers.background import BackgroundScheduler
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PIL import Image
from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from .config import get_settings
from .database import Base, SessionLocal, engine, get_db
from .deps import admin_user, current_user, owns
from .models import *
from .realtime import hub
from .schemas import *
from .security import create_access_token, hash_password, hash_token, new_participation_token, verify_password
from .services import aware, cleanup_expired, finalize_votes, results, vote_count

settings = get_settings()
vote_attempts: dict[str, deque[datetime]] = defaultdict(deque)


def bootstrap():
    Base.metadata.create_all(engine)
    with SessionLocal.begin() as db:
        if settings.admin_username and settings.admin_password:
            existing = db.scalar(select(User).where(User.username == settings.admin_username))
            if not existing:
                db.add(User(username=settings.admin_username, full_name="Administrator", password_hash=hash_password(settings.admin_password), role=Role.ADMIN))
        if not db.get(GlobalBranding, 1):
            db.add(GlobalBranding(id=1, university_name="", school_name="", background_type="none", background_opacity=Decimal("0.12")))


@asynccontextmanager
async def lifespan(app):
    bootstrap()
    scheduler = BackgroundScheduler(timezone="UTC")
    def scheduled_cleanup():
        with SessionLocal() as db:
            cleanup_expired(db)
    scheduler.add_job(scheduled_cleanup, "interval", minutes=settings.cleanup_interval_minutes, id="token-cleanup", max_instances=1)
    scheduler.start()
    yield
    scheduler.shutdown(wait=False)


app = FastAPI(title=settings.app_name, version="2.2", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=settings.origins, allow_credentials=False, allow_methods=["*"], allow_headers=["*"])
app.mount("/assets", StaticFiles(directory=settings.asset_dir), name="assets")


def user_json(u: User):
    return {"id": u.id, "username": u.username, "full_name": u.full_name, "role": u.role.value, "active": u.active}


def session_for_user(db, sid, user):
    s = db.get(PresentationSession, sid)
    if not s:
        raise HTTPException(404, "Δεν βρέθηκε")
    owns(s.owner_id, user)
    return s


def current_open(db, session_id):
    return db.scalar(select(Presentation).where(Presentation.session_id == session_id, Presentation.status == PresentationStatus.VOTING_OPEN))


def branding_json(b):
    return {k: getattr(b, k) for k in ("university_name", "school_name", "department_name", "background_type", "background_value")} | {
        "logo_url": f"/assets/{b.logo_path}" if b.logo_path else None,
        "background_image_url": f"/assets/{b.background_image_path}" if b.background_image_path else None,
        "background_opacity": float(b.background_opacity or 0.12),
    }


def effective_branding(db, course):
    global_b = db.get(GlobalBranding, 1)
    data = branding_json(global_b)
    if course:
        course_data = branding_json(course)
        for key, value in course_data.items():
            if value not in (None, ""):
                data[key] = value
    return data


def public_state(db, s):
    course, period, group = db.get(Course, s.course_id), db.get(AcademicPeriod, s.period_id), db.get(StudentGroup, s.group_id)
    p = current_open(db, s.id)
    criteria = db.scalars(select(Criterion).where(Criterion.session_id == s.id).order_by(Criterion.position)).all()
    return {"public_id": s.public_id, "title": s.title, "course": course.name, "period": period.name, "group": group.title,
            "session_date": s.session_date.isoformat(), "status": s.status.value, "lock_new_participants": s.lock_new_participants,
            "results_revealed": s.results_revealed, "active_presentation": None if not p else {"id": p.id, "presenter_name": p.presenter_name, "title": p.title, "closes_at": p.voting_closes_at},
            "criteria": [{"id": c.id, "name": c.name, "weight": c.weight} for c in criteria],
            "vote_count": vote_count(db, p.id) if p else 0, "participation_url": f"{settings.public_base_url}/join/{s.public_id}",
            "branding": effective_branding(db, course), "results": results(db, s.id) if s.results_revealed else None}


def check_vote_rate(token: str):
    """Limit bursts by participant credential, never by network address."""
    now = datetime.now(timezone.utc); attempts = vote_attempts[hash_token(token)]
    while attempts and (now - attempts[0]).total_seconds() > 10: attempts.popleft()
    if len(attempts) >= 8: raise HTTPException(429, "Πάρα πολλές υποβολές· περιμένετε λίγο")
    attempts.append(now)


@app.get("/api/health")
def health(): return {"status": "ok", "version": "2.2"}


@app.post("/api/auth/login")
def login(data: LoginIn, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == data.username))
    if not user or not user.active or not verify_password(data.password, user.password_hash):
        raise HTTPException(401, "Λανθασμένα στοιχεία σύνδεσης")
    return {"access_token": create_access_token(user.id, user.role.value), "token_type": "bearer", "user": user_json(user)}


@app.get("/api/auth/me")
def me(user=Depends(current_user)): return user_json(user)


@app.put("/api/auth/password")
def password(data: dict, user=Depends(current_user), db: Session = Depends(get_db)):
    if not verify_password(data.get("current_password", ""), user.password_hash) or len(data.get("new_password", "")) < 10:
        raise HTTPException(400, "Μη έγκυρος κωδικός")
    user.password_hash = hash_password(data["new_password"]); db.commit(); return {"ok": True}


@app.get("/api/users")
def users(_=Depends(admin_user), db: Session = Depends(get_db)): return [user_json(x) for x in db.scalars(select(User)).all()]


@app.post("/api/users", status_code=201)
def create_user(data: UserCreate, _=Depends(admin_user), db: Session = Depends(get_db)):
    u = User(username=data.username, full_name=data.full_name, password_hash=hash_password(data.password), role=data.role)
    db.add(u)
    try: db.commit()
    except IntegrityError: db.rollback(); raise HTTPException(409, "Το όνομα χρήστη υπάρχει ήδη")
    return user_json(u)


@app.patch("/api/users/{uid}")
def update_user(uid: int, data: UserUpdate, admin=Depends(admin_user), db: Session = Depends(get_db)):
    u = db.get(User, uid)
    if not u: raise HTTPException(404)
    for key, value in data.model_dump(exclude_unset=True).items():
        setattr(u, "password_hash" if key == "password" else key, hash_password(value) if key == "password" else value)
    if u.id == admin.id and not u.active: raise HTTPException(400, "Δεν μπορείτε να απενεργοποιήσετε τον εαυτό σας")
    db.commit(); return user_json(u)


@app.get("/api/courses")
def list_courses(user=Depends(current_user), db: Session = Depends(get_db)):
    q = select(Course) if user.role == Role.ADMIN else select(Course).where(Course.owner_id == user.id)
    return [{"id": x.id, "name": x.name, "description": x.description, "owner_id": x.owner_id, "branding": effective_branding(db, x)} for x in db.scalars(q).all()]


@app.post("/api/courses", status_code=201)
def create_course(data: CourseIn, user=Depends(current_user), db: Session = Depends(get_db)):
    item = Course(**data.model_dump(), owner_id=user.id); db.add(item); db.commit(); db.refresh(item); return {"id": item.id, **data.model_dump()}


@app.get("/api/periods")
def periods(user=Depends(current_user), db: Session = Depends(get_db)):
    q = select(AcademicPeriod) if user.role == Role.ADMIN else select(AcademicPeriod).where(AcademicPeriod.owner_id == user.id)
    return [{"id": x.id, "name": x.name} for x in db.scalars(q).all()]


@app.post("/api/periods", status_code=201)
def create_period(data: PeriodIn, user=Depends(current_user), db: Session = Depends(get_db)):
    x = AcademicPeriod(name=data.name, owner_id=user.id); db.add(x); db.commit(); return {"id": x.id, "name": x.name}


@app.get("/api/groups")
def groups(user=Depends(current_user), db: Session = Depends(get_db)):
    q = select(StudentGroup) if user.role == Role.ADMIN else select(StudentGroup).where(StudentGroup.owner_id == user.id)
    return [{"id": x.id, "title": x.title, "course_id": x.course_id, "period_id": x.period_id} for x in db.scalars(q).all()]


@app.post("/api/groups", status_code=201)
def create_group(data: GroupIn, user=Depends(current_user), db: Session = Depends(get_db)):
    course, period = db.get(Course, data.course_id), db.get(AcademicPeriod, data.period_id)
    if not course or not period: raise HTTPException(400, "Μη έγκυρη επιλογή")
    owns(course.owner_id, user); owns(period.owner_id, user)
    x = StudentGroup(**data.model_dump(), owner_id=user.id); db.add(x); db.commit(); return {"id": x.id, **data.model_dump()}


@app.get("/api/groups/{gid}/students")
def students(gid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    g = db.get(StudentGroup, gid); owns(g.owner_id, user) if g else (_ for _ in ()).throw(HTTPException(404))
    return [{"id": x.id, "full_name": x.full_name, "presentation_title": x.presentation_title} for x in db.scalars(select(Student).where(Student.group_id == gid)).all()]


@app.post("/api/groups/{gid}/students")
def add_student(gid: int, data: StudentIn, user=Depends(current_user), db: Session = Depends(get_db)):
    g = db.get(StudentGroup, gid); owns(g.owner_id, user) if g else (_ for _ in ()).throw(HTTPException(404))
    x = Student(group_id=gid, **data.model_dump()); db.add(x); db.commit(); return {"id": x.id, **data.model_dump()}


def parse_import(file: UploadFile):
    raw = file.file.read(settings.upload_max_bytes + 1)
    if len(raw) > settings.upload_max_bytes: raise HTTPException(413, "Το αρχείο είναι πολύ μεγάλο")
    rows = []
    try:
        if file.filename and file.filename.lower().endswith(".csv"):
            text = raw.decode("utf-8-sig"); reader = csv.DictReader(io.StringIO(text))
            for r in reader: rows.append({"full_name": (r.get("Student Name") or r.get("Ονοματεπώνυμο") or "").strip(), "presentation_title": (r.get("Presentation Title") or r.get("Τίτλος Παρουσίασης") or "").strip() or None})
        elif file.filename and file.filename.lower().endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True); ws = wb.active
            headers = [str(x.value or "").strip() for x in next(ws.iter_rows())]
            for values in ws.iter_rows(values_only=True):
                r = dict(zip(headers, values)); rows.append({"full_name": str(r.get("Student Name") or r.get("Ονοματεπώνυμο") or "").strip(), "presentation_title": str(r.get("Presentation Title") or r.get("Τίτλος Παρουσίασης") or "").strip() or None})
        else: raise HTTPException(415, "Υποστηρίζονται CSV και XLSX")
    except HTTPException: raise
    except Exception: raise HTTPException(400, "Δεν ήταν δυνατή η ανάγνωση του αρχείου")
    rows = [r for r in rows if r["full_name"]]
    counts = {n: sum(x["full_name"] == n for x in rows) for n in {x["full_name"] for x in rows}}
    return rows, [n for n, c in counts.items() if c > 1]


@app.post("/api/groups/{gid}/import/preview")
def preview_import(gid: int, file: UploadFile = File(...), user=Depends(current_user), db: Session = Depends(get_db)):
    g = db.get(StudentGroup, gid); owns(g.owner_id, user) if g else (_ for _ in ()).throw(HTTPException(404))
    rows, warnings = parse_import(file); return {"rows": rows, "duplicate_names": warnings}


@app.post("/api/groups/{gid}/import")
def commit_import(gid: int, file: UploadFile = File(...), user=Depends(current_user), db: Session = Depends(get_db)):
    g = db.get(StudentGroup, gid); owns(g.owner_id, user) if g else (_ for _ in ()).throw(HTTPException(404))
    rows, warnings = parse_import(file)
    for row in rows: db.add(Student(group_id=gid, **row))
    db.commit(); return {"imported": len(rows), "duplicate_names": warnings}


@app.get("/api/import-template.{kind}")
def template(kind: str, _=Depends(current_user)):
    if kind == "csv": return Response("Student Name,Presentation Title\nΜαρία Παπαδοπούλου,Παράδειγμα\n", media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=presenters.csv"})
    if kind == "xlsx":
        wb = Workbook(); ws = wb.active; ws.append(["Student Name", "Presentation Title"]); ws.append(["Μαρία Παπαδοπούλου", "Παράδειγμα"]); out=io.BytesIO(); wb.save(out)
        return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=presenters.xlsx"})
    raise HTTPException(404)


@app.get("/api/sessions")
def sessions(user=Depends(current_user), db: Session = Depends(get_db)):
    q = select(PresentationSession) if user.role == Role.ADMIN else select(PresentationSession).where(PresentationSession.owner_id == user.id)
    return [{"id": x.id, "public_id": x.public_id, "title": x.title, "session_date": x.session_date, "status": x.status.value, "is_demo": x.is_demo, "course_id": x.course_id, "group_id": x.group_id} for x in db.scalars(q.order_by(PresentationSession.session_date.desc())).all()]


@app.post("/api/sessions", status_code=201)
def create_session(data: SessionCreate, user=Depends(current_user), db: Session = Depends(get_db)):
    course, period, group = db.get(Course, data.course_id), db.get(AcademicPeriod, data.period_id), db.get(StudentGroup, data.group_id)
    if not all((course, period, group)): raise HTTPException(400, "Μη έγκυρη επιλογή")
    for x in (course, period, group): owns(x.owner_id, user)
    s = PresentationSession(public_id=secrets.token_urlsafe(12), owner_id=user.id, **data.model_dump(exclude={"criteria", "presenters"})); db.add(s); db.flush()
    for i, c in enumerate(data.criteria): db.add(Criterion(session_id=s.id, position=i, **c.model_dump()))
    presenters = data.presenters or [StudentIn(full_name=x.full_name, presentation_title=x.presentation_title) for x in db.scalars(select(Student).where(Student.group_id == data.group_id)).all()]
    for i, p in enumerate(presenters): db.add(Presentation(session_id=s.id, presenter_name=p.full_name, title=p.presentation_title, position=i))
    db.commit(); return {"id": s.id, "public_id": s.public_id}


@app.get("/api/sessions/{sid}")
def session_detail(sid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user); ps = db.scalars(select(Presentation).where(Presentation.session_id == sid).order_by(Presentation.position)).all(); cs = db.scalars(select(Criterion).where(Criterion.session_id == sid).order_by(Criterion.position)).all()
    return {"id": s.id, "public_id": s.public_id, "title": s.title, "status": s.status.value, "session_date": s.session_date, "lock_new_participants": s.lock_new_participants, "is_demo": s.is_demo, "results_revealed": s.results_revealed,
            "presentations": [{"id": p.id, "presenter_name": p.presenter_name, "title": p.title, "status": p.status.value, "vote_count": vote_count(db, p.id, s.status == SessionStatus.COMPLETED)} for p in ps],
            "criteria": [{"id": c.id, "name": c.name, "weight": c.weight} for c in cs], "results": results(db, sid) if s.status == SessionStatus.COMPLETED else None,
            "participation_url": f"{settings.public_base_url}/join/{s.public_id}"}


@app.put("/api/sessions/{sid}/criteria")
def set_criteria(sid: int, data: list[CriterionIn], user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user)
    if s.status != SessionStatus.DRAFT or s.criteria_locked: raise HTTPException(409, "Τα κριτήρια έχουν κλειδωθεί")
    if not data or sum(x.weight for x in data) != 100: raise HTTPException(422, "Τα βάρη πρέπει να αθροίζουν σε 100%")
    db.execute(delete(Criterion).where(Criterion.session_id == sid))
    for position, item in enumerate(data): db.add(Criterion(session_id=sid, position=position, **item.model_dump()))
    db.commit(); return {"ok": True}


@app.post("/api/sessions/{sid}/activate")
async def activate(sid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user)
    if s.status != SessionStatus.DRAFT: raise HTTPException(409, "Η συνεδρία δεν είναι πρόχειρη")
    if db.scalar(select(func.sum(Criterion.weight)).where(Criterion.session_id == sid)) != 100: raise HTTPException(400, "Τα βάρη πρέπει να αθροίζουν σε 100%")
    s.status = SessionStatus.ACTIVE; db.commit(); await hub.broadcast(s.public_id, "state_changed"); return {"ok": True}


@app.patch("/api/sessions/{sid}/admission")
async def admission(sid: int, data: dict, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user)
    if s.status != SessionStatus.ACTIVE: raise HTTPException(409)
    s.lock_new_participants = bool(data.get("locked")); db.commit(); await hub.broadcast(s.public_id, "state_changed"); return {"locked": s.lock_new_participants}


@app.post("/api/sessions/{sid}/presentations/{pid}/open")
async def open_voting(sid: int, pid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user); p = db.get(Presentation, pid)
    if s.status != SessionStatus.ACTIVE or not p or p.session_id != sid or p.status != PresentationStatus.PENDING: raise HTTPException(409, "Δεν μπορεί να ανοίξει η ψηφοφορία")
    if current_open(db, sid): raise HTTPException(409, "Υπάρχει ήδη ανοικτή ψηφοφορία")
    now = datetime.now(timezone.utc); p.status = PresentationStatus.VOTING_OPEN; p.voting_opened_at = now; p.voting_closes_at = now + timedelta(seconds=s.voting_duration); s.criteria_locked = True; db.commit()
    await hub.broadcast(s.public_id, "state_changed"); return {"ok": True}


@app.post("/api/sessions/{sid}/presentations/{pid}/close")
async def close_voting(sid: int, pid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user); p = db.get(Presentation, pid)
    if not p or p.session_id != sid or p.status != PresentationStatus.VOTING_OPEN: raise HTTPException(409)
    p.status = PresentationStatus.EVALUATED if vote_count(db, pid) else PresentationStatus.NO_VOTES; p.voting_closed_at = datetime.now(timezone.utc); db.commit(); await hub.broadcast(s.public_id, "state_changed"); return {"ok": True}


@app.post("/api/sessions/{sid}/presentations/{pid}/extend")
async def extend_voting(sid: int, pid: int, data: dict, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user); p = db.get(Presentation, pid); seconds = int(data.get("seconds", 60))
    if not p or p.session_id != sid or p.status != PresentationStatus.VOTING_OPEN or not 1 <= seconds <= 3600: raise HTTPException(409)
    p.voting_closes_at = max(aware(p.voting_closes_at), datetime.now(timezone.utc)) + timedelta(seconds=seconds); db.commit()
    await hub.broadcast(s.public_id, "state_changed"); return {"closes_at": p.voting_closes_at}


@app.post("/api/sessions/{sid}/presentations/{pid}/skip")
async def skip(sid: int, pid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    s=session_for_user(db,sid,user); p=db.get(Presentation,pid)
    if not p or p.session_id != sid or p.status not in (PresentationStatus.PENDING, PresentationStatus.SKIPPED): raise HTTPException(409)
    p.status = PresentationStatus.PENDING if p.status == PresentationStatus.SKIPPED else PresentationStatus.SKIPPED; db.commit(); await hub.broadcast(s.public_id,"state_changed"); return {"status":p.status.value}


@app.put("/api/sessions/{sid}/presentations/{pid}")
async def edit_presentation(sid: int, pid: int, data: PresentationEditIn, user=Depends(current_user), db: Session = Depends(get_db)):
    s=session_for_user(db,sid,user); p=db.get(Presentation,pid)
    if not p or p.session_id != sid or s.status in (SessionStatus.COMPLETED,SessionStatus.ARCHIVED): raise HTTPException(409)
    changes=(("presenter_name",p.presenter_name,data.presenter_name),("title",p.title,data.title))
    for field, old, new in changes:
        if old != new: db.add(PresentationEdit(presentation_id=p.id,field=field,previous_value=old,new_value=new,changed_by_id=user.id)); setattr(p,field,new)
    db.commit(); await hub.broadcast(s.public_id,"presenter_updated"); return {"id":p.id,"presenter_name":p.presenter_name,"title":p.title}


@app.get("/api/sessions/{sid}/presentations/{pid}/history")
def edit_history(sid:int,pid:int,user=Depends(current_user),db:Session=Depends(get_db)):
    session_for_user(db,sid,user); p=db.get(Presentation,pid)
    if not p or p.session_id != sid: raise HTTPException(404)
    return [{"field":x.field,"previous_value":x.previous_value,"new_value":x.new_value,"changed_by_id":x.changed_by_id,"changed_at":x.changed_at} for x in db.scalars(select(PresentationEdit).where(PresentationEdit.presentation_id==pid).order_by(PresentationEdit.changed_at.desc())).all()]


@app.post("/api/sessions/{sid}/complete")
async def complete(sid:int,user=Depends(current_user),db:Session=Depends(get_db)):
    s=session_for_user(db,sid,user)
    if s.status != SessionStatus.ACTIVE: raise HTTPException(409)
    p=current_open(db,sid)
    if p: p.status=PresentationStatus.EVALUATED if vote_count(db,p.id) else PresentationStatus.NO_VOTES; p.voting_closed_at=datetime.now(timezone.utc)
    for pending in db.scalars(select(Presentation).where(Presentation.session_id==sid,Presentation.status==PresentationStatus.PENDING)).all(): pending.status=PresentationStatus.SKIPPED
    finalize_votes(db,sid); s.status=SessionStatus.COMPLETED; s.completed_at=datetime.now(timezone.utc); db.commit(); await hub.broadcast(s.public_id,"state_changed"); return {"ok":True}


@app.post("/api/sessions/{sid}/reveal")
async def reveal(sid:int,user=Depends(current_user),db:Session=Depends(get_db)):
    s=session_for_user(db,sid,user)
    if s.status != SessionStatus.COMPLETED: raise HTTPException(409)
    s.results_revealed=True; db.commit(); await hub.broadcast(s.public_id,"state_changed"); return {"ok":True}


@app.post("/api/sessions/{sid}/duplicate")
def duplicate(sid:int,data:DuplicateIn,user=Depends(current_user),db:Session=Depends(get_db)):
    old=session_for_user(db,sid,user); gid=data.group_id or old.group_id; group=db.get(StudentGroup,gid); owns(group.owner_id,user)
    new=PresentationSession(public_id=secrets.token_urlsafe(12),course_id=old.course_id,period_id=old.period_id,group_id=gid,owner_id=user.id,session_date=data.session_date,title=old.title,status=SessionStatus.DRAFT,voting_duration=old.voting_duration,is_demo=old.is_demo); db.add(new); db.flush()
    for c in db.scalars(select(Criterion).where(Criterion.session_id==sid)).all(): db.add(Criterion(session_id=new.id,name=c.name,weight=c.weight,position=c.position))
    source=db.scalars(select(Student).where(Student.group_id==gid)).all()
    for i,x in enumerate(source): db.add(Presentation(session_id=new.id,student_id=x.id,presenter_name=x.full_name,title=x.presentation_title,position=i))
    db.commit(); return {"id":new.id,"public_id":new.public_id}


@app.post("/api/sessions/{sid}/archive")
def archive(sid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user)
    if s.status != SessionStatus.COMPLETED: raise HTTPException(409)
    s.status = SessionStatus.ARCHIVED; db.commit(); return {"ok": True}


@app.patch("/api/sessions/{sid}/owner")
def transfer_session(sid: int, data: dict, _=Depends(admin_user), db: Session = Depends(get_db)):
    s = db.get(PresentationSession, sid); new_owner = db.get(User, int(data.get("owner_id", 0)))
    if not s or not new_owner: raise HTTPException(400)
    s.owner_id = new_owner.id; db.commit(); return {"ok": True}


@app.post("/api/sessions/{sid}/demo/simulate")
def simulate_demo(sid: int, data: dict, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user)
    if not s.is_demo or s.status != SessionStatus.DRAFT: raise HTTPException(409, "Απαιτείται πρόχειρη DEMO συνεδρία")
    count = min(max(int(data.get("participants", 10)), 1), 500); criteria = db.scalars(select(Criterion).where(Criterion.session_id == sid)).all(); presentations = db.scalars(select(Presentation).where(Presentation.session_id == sid)).all()
    for p in presentations:
        p.status = PresentationStatus.EVALUATED
        for _ in range(count):
            anonymous_id = uuid4().hex
            for c in criteria: db.add(AnonymousVoteScore(presentation_id=p.id, anonymous_vote_id=anonymous_id, criterion_id=c.id, score=random.randint(1, 5), synthetic=True))
    s.status = SessionStatus.COMPLETED; s.completed_at = datetime.now(timezone.utc); db.commit(); return {"generated_votes_per_presentation": count}


@app.delete("/api/sessions/{sid}/demo", status_code=204)
def delete_demo(sid: int, user=Depends(current_user), db: Session = Depends(get_db)):
    s = session_for_user(db, sid, user)
    if not s.is_demo: raise HTTPException(409)
    db.delete(s); db.commit(); return Response(status_code=204)


@app.get("/api/public/sessions/{public_id}")
def get_public(public_id:str,db:Session=Depends(get_db)):
    s=db.scalar(select(PresentationSession).where(PresentationSession.public_id==public_id))
    if not s: raise HTTPException(404)
    return public_state(db,s)


@app.post("/api/public/sessions/{public_id}/tokens")
def issue_token(public_id:str,db:Session=Depends(get_db)):
    s=db.scalar(select(PresentationSession).where(PresentationSession.public_id==public_id))
    if not s or s.status != SessionStatus.ACTIVE: raise HTTPException(409,"Η συνεδρία δεν είναι ενεργή")
    if s.lock_new_participants: raise HTTPException(403,"Η είσοδος νέων συμμετεχόντων είναι κλειδωμένη")
    raw=new_participation_token(); now=datetime.now(timezone.utc); db.add(ParticipationToken(session_id=s.id,token_hash=hash_token(raw),issued_at=now,expires_at=now+timedelta(hours=settings.token_lifetime_hours))); db.commit()
    return {"token":raw,"expires_at":now+timedelta(hours=settings.token_lifetime_hours)}


def participant(db,public_id,token):
    s=db.scalar(select(PresentationSession).where(PresentationSession.public_id==public_id)); t=db.scalar(select(ParticipationToken).where(ParticipationToken.token_hash==hash_token(token or ""),ParticipationToken.session_id==s.id)) if s else None
    if not s or s.status != SessionStatus.ACTIVE or not t or t.revoked or aware(t.expires_at)<=datetime.now(timezone.utc): raise HTTPException(401,"Το διακριτικό συμμετοχής έληξε")
    return s,t


@app.get("/api/public/sessions/{public_id}/my-vote")
def my_vote(public_id:str,x_participation_token:str|None=Header(None),db:Session=Depends(get_db)):
    s,t=participant(db,public_id,x_participation_token); p=current_open(db,s.id)
    if not p:return {"scores":{}}
    v=db.scalar(select(Vote).where(Vote.presentation_id==p.id,Vote.token_id==t.id))
    return {"scores":{} if not v else {str(x.criterion_id):x.score for x in db.scalars(select(VoteScore).where(VoteScore.vote_id==v.id)).all()}}


@app.put("/api/public/sessions/{public_id}/vote")
async def submit_vote(public_id:str,data:VoteIn,x_participation_token:str|None=Header(None),db:Session=Depends(get_db)):
    check_vote_rate(x_participation_token or "")
    s,t=participant(db,public_id,x_participation_token); p=current_open(db,s.id)
    if not p or (p.voting_closes_at and aware(p.voting_closes_at)<=datetime.now(timezone.utc)): raise HTTPException(409,"Η ψηφοφορία έχει κλείσει")
    criteria=db.scalars(select(Criterion).where(Criterion.session_id==s.id)).all(); ids={c.id for c in criteria}
    if set(data.scores)!=ids or any(v<1 or v>5 for v in data.scores.values()): raise HTTPException(422,"Απαιτείται βαθμός 1–5 για κάθε κριτήριο")
    try:
        v=db.scalar(select(Vote).where(Vote.presentation_id==p.id,Vote.token_id==t.id))
        if not v: v=Vote(presentation_id=p.id,token_id=t.id); db.add(v); db.flush()
        existing={x.criterion_id:x for x in db.scalars(select(VoteScore).where(VoteScore.vote_id==v.id)).all()}
        for cid,score in data.scores.items():
            if cid in existing: existing[cid].score=score
            else: db.add(VoteScore(vote_id=v.id,criterion_id=cid,score=score))
        db.commit()
    except IntegrityError: db.rollback(); raise HTTPException(409,"Ταυτόχρονη υποβολή· δοκιμάστε ξανά")
    await hub.broadcast(s.public_id,"vote_count_changed"); return {"ok":True,"vote_count":vote_count(db,p.id)}


@app.websocket("/api/ws/{public_id}")
async def websocket(public_id:str,ws:WebSocket):
    with SessionLocal() as db:
        if not db.scalar(select(PresentationSession).where(PresentationSession.public_id==public_id)): await ws.close(code=1008); return
    await hub.connect(public_id,ws)
    try:
        while True: await ws.receive_text()
    except WebSocketDisconnect: hub.disconnect(public_id,ws)


def apply_branding(target,data):
    for key,value in data.model_dump().items(): setattr(target,key,Decimal(str(value)) if key=="background_opacity" and value is not None else value)


@app.get("/api/branding")
def branding(_=Depends(admin_user),db:Session=Depends(get_db)): return branding_json(db.get(GlobalBranding,1))


@app.get("/api/public/branding")
def public_branding(db:Session=Depends(get_db)): return branding_json(db.get(GlobalBranding,1))


@app.put("/api/branding")
def set_branding(data:BrandingIn,_=Depends(admin_user),db:Session=Depends(get_db)):
    b=db.get(GlobalBranding,1); apply_branding(b,data); db.commit(); return branding_json(b)


@app.put("/api/courses/{cid}/branding")
def course_branding(cid:int,data:BrandingIn,user=Depends(current_user),db:Session=Depends(get_db)):
    c=db.get(Course,cid); owns(c.owner_id,user) if c else (_ for _ in ()).throw(HTTPException(404)); apply_branding(c,data); db.commit(); return effective_branding(db,c)


def safe_image(upload:UploadFile):
    raw=upload.file.read(settings.upload_max_bytes+1)
    if len(raw)>settings.upload_max_bytes: raise HTTPException(413)
    try:
        image=Image.open(io.BytesIO(raw)); image.verify(); fmt=image.format
    except Exception: raise HTTPException(415,"Μη έγκυρη εικόνα")
    if fmt not in {"PNG","JPEG","WEBP"}: raise HTTPException(415,"Υποστηρίζονται PNG, JPEG και WebP. Το SVG απορρίπτεται για ασφάλεια.")
    ext={"PNG":"png","JPEG":"jpg","WEBP":"webp"}[fmt]; name=f"{uuid4().hex}.{ext}"; Path(settings.asset_dir,name).write_bytes(raw); return name


def spreadsheet_safe(value):
    """Neutralize spreadsheet formulas while preserving Unicode text."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@app.post("/api/branding/logo")
def upload_global_logo(file:UploadFile=File(...),_=Depends(admin_user),db:Session=Depends(get_db)):
    b=db.get(GlobalBranding,1); b.logo_path=safe_image(file); db.commit(); return branding_json(b)


@app.post("/api/courses/{cid}/branding/logo")
def upload_course_logo(cid:int,file:UploadFile=File(...),user=Depends(current_user),db:Session=Depends(get_db)):
    c=db.get(Course,cid); owns(c.owner_id,user) if c else (_ for _ in ()).throw(HTTPException(404)); c.logo_path=safe_image(file); db.commit(); return effective_branding(db,c)


@app.post("/api/branding/background")
def upload_global_background(file:UploadFile=File(...),_=Depends(admin_user),db:Session=Depends(get_db)):
    b=db.get(GlobalBranding,1); b.background_image_path=safe_image(file); b.background_type="image"; db.commit(); return branding_json(b)


@app.post("/api/courses/{cid}/branding/background")
def upload_course_background(cid:int,file:UploadFile=File(...),user=Depends(current_user),db:Session=Depends(get_db)):
    c=db.get(Course,cid); owns(c.owner_id,user) if c else (_ for _ in ()).throw(HTTPException(404)); c.background_image_path=safe_image(file); c.background_type="image"; db.commit(); return effective_branding(db,c)


def export_rows(db,s):
    course,period,group=db.get(Course,s.course_id),db.get(AcademicPeriod,s.period_id),db.get(StudentGroup,s.group_id)
    rows=[]
    for item in results(db,s.id):
        base={"Course":course.name,"Academic Period":period.name,"Group":group.title,"Session Date":s.session_date.isoformat(),"Presenter":item["presenter_name"],"Presentation Title":item["title"] or "","Status":item["status"],"Valid Votes":item["vote_count"],"Weighted Score":item["weighted_score"],"Rank":item["rank"]}
        for c in item["criteria"]: base[f'{c["name"]} ({c["weight"]}%)']=c["mean"]
        rows.append({key: spreadsheet_safe(value) for key, value in base.items()})
    return rows,course


@app.get("/api/sessions/{sid}/export.{kind}")
def export(sid:int,kind:str,user=Depends(current_user),db:Session=Depends(get_db)):
    s=session_for_user(db,sid,user); rows,course=export_rows(db,s)
    if kind=="csv":
        out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=list(rows[0]) if rows else ["Course"]); writer.writeheader(); writer.writerows(rows)
        return Response("\ufeff"+out.getvalue(),media_type="text/csv; charset=utf-8",headers={"Content-Disposition":f"attachment; filename=session-{sid}.csv"})
    if kind=="xlsx":
        wb=Workbook(); summary=wb.active; summary.title="Summary"; brand=effective_branding(db,course); summary.append([spreadsheet_safe(brand.get("university_name") or "")]); summary.append([spreadsheet_safe(brand.get("school_name") or "")]); summary.append([spreadsheet_safe(brand.get("department_name") or "")]); summary.append([])
        if rows: summary.append(list(rows[0])); [summary.append(list(r.values())) for r in rows]; summary[1][0].font=Font(bold=True,size=14)
        detail=wb.create_sheet("Anonymous Scores"); detail.append(["Presentation ID","Anonymous Vote","Criterion","Score","Synthetic"])
        for x in db.scalars(select(AnonymousVoteScore).join(Presentation).where(Presentation.session_id==sid)).all(): detail.append([x.presentation_id,x.anonymous_vote_id,x.criterion_id,x.score,x.synthetic])
        out=io.BytesIO(); wb.save(out); return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f"attachment; filename=session-{sid}.xlsx"})
    raise HTTPException(404)
